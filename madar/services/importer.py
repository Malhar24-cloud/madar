"""استيراد الموظفين من Excel أو CSV مع فحص كل صف قبل الحفظ.
الخطوة الأولى تفحص وتعرض الأخطاء، ولا يُحفظ شيء إلا بعد موافقة مدير الموارد البشرية على النتيجة."""
import csv
import io
import re
import secrets
from datetime import date, datetime

from flask import current_app

from .. import policy
from ..models import Employee, User, db
from . import audit, mailer, tokens

COLUMNS = [
    ("code", "الرقم الوظيفي", False, "اتركه فاضي ويولّده النظام"),
    ("name", "الاسم الكامل", True, ""),
    ("national_id", "رقم الهوية أو الإقامة", True, "10 أرقام"),
    ("nationality", "الجنسية", True, "سعودي، هندي، ..."),
    ("department", "القسم", True, " / ".join(policy.DEPARTMENTS)),
    ("title", "المسمى الوظيفي", True, ""),
    ("manager_code", "الرقم الوظيفي للمدير", False, "رقم موجود في الملف أو في النظام"),
    ("hire_date", "تاريخ المباشرة", True, "YYYY-MM-DD"),
    ("basic", "الراتب الأساسي", True, ""),
    ("housing", "بدل السكن", False, "فاضي = 25% من الأساسي"),
    ("transport", "بدل النقل", False, "فاضي = 10% من الأساسي"),
    ("contract_type", "نوع العقد", False, "محدد أو غير محدد"),
    ("end_date", "تاريخ انتهاء العقد", False, "إلزامي للعقد المحدد"),
    ("leave_balance", "رصيد الإجازات الحالي", False, "بالأيام"),
    ("bank", "البنك", False, ""),
    ("iban", "الآيبان", False, "SA + 22 خانة"),
    ("email", "البريد الوظيفي", False, "فاضي للعمال بدون بريد: يأخذون رمز تفعيل"),
    ("phone", "الجوال", False, "05xxxxxxxx"),
    ("iqama_expiry", "انتهاء الإقامة", False, "لغير السعوديين"),
    ("gosi_new_system", "مسجل في التأمينات بعد 2024-07-03", False, "نعم أو لا"),
    ("attendance_id", "رقم البصمة", False, "الرقم في جهاز الحضور"),
    ("role", "الدور في النظام", False, "employee / manager (الأدوار الأعلى من صفحة الصلاحيات أو بواسطة مدير الموارد البشرية)"),
]
KEYS = [c[0] for c in COLUMNS]
ROLE_VALUES = {"employee", "manager", "admin", "payroll", "recruit"}


def template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "employees"
    ws.sheet_view.rightToLeft = True
    ws.append([c[1] for c in COLUMNS])
    ws.append([c[0] for c in COLUMNS])
    ws.append([c[3] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6A61")
    for cell in ws[3]:
        cell.font = Font(italic=True, color="74858A")
    ws.append(["", "مثال: محمد أحمد", "1012345678", "سعودي", "العمليات", "مشرف مواقع", "", "2024-03-01", 8500, "", "",
               "غير محدد", "", 12, "الراجحي", "SA0380000000608010167519", "", "0501234567", "", "لا", "15", "employee"])
    for i, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(c[1]) + 4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


MAX_ROWS = 3000


def read_rows(fs):
    name = (fs.filename or "").lower()
    data = fs.read()
    if name.endswith(".xlsx"):
        from ..filesafe import check_zip
        check_zip(data)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = []
        for r in ws.iter_rows(values_only=True, max_col=40):
            rows.append(list(r))
            if len(rows) > MAX_ROWS + 3:
                raise ValueError(f"الملف أكبر من الحد ({MAX_ROWS} صف). قسّمه على أكثر من ملف.")
    elif name.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
        if len(rows) > MAX_ROWS + 3:
            raise ValueError(f"الملف أكبر من الحد ({MAX_ROWS} صف). قسّمه على أكثر من ملف.")
    else:
        raise ValueError("ارفع ملف Excel (xlsx) أو CSV من القالب.")
    rows = [r for r in rows if r and any(str(c or "").strip() for c in r)]
    # نتعرف على صف المفاتيح الإنجليزية (الصف الثاني في القالب) أو نعتمد ترتيب القالب
    key_row = next((i for i, r in enumerate(rows[:3]) if "name" in [str(c or "").strip() for c in r]), None)
    if key_row is not None:
        keys = [str(c or "").strip() for c in rows[key_row]]
        body = rows[key_row + 1:]
    else:
        keys, body = KEYS, rows[1:]
    out = []
    for r in body:
        rec = {k: r[i] if i < len(r) else None for i, k in enumerate(keys) if k in KEYS}
        if str(rec.get("name") or "").startswith("مثال") or str(rec.get("code") or "").startswith("اتركه"):
            continue
        out.append(rec)
    return out


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _s(v)
    if not s:
        return None
    for f in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    raise ValueError(s)


def _num(v):
    from ..validate import MAX_MONEY, number
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = repr(v)
    n = number(_s(v), "الرقم", 0, MAX_MONEY, allow_empty=True)
    return None if n is None else int(round(n))


def validate(records):
    from .. import validate as vd
    existing_att = {str(e.attendance_id) for e in Employee.query.all() if e.attendance_id}
    existing_codes = {e.code for e in Employee.query.all()}
    existing_emails = {u.email for u in User.query.all() if u.email} | {e.email for e in Employee.query.all() if e.email}
    existing_ids = {e.national_id for e in Employee.query.all() if e.national_id}
    file_codes = {_s(r.get("code")).upper() for r in records if _s(r.get("code"))}
    seen = {"code": set(), "email": set(), "national_id": set(), "attendance_id": set()}
    results = []
    for n, r in enumerate(records, start=1):
        errs, clean = [], {}
        for key, label, required, _ in COLUMNS:
            if required and not _s(r.get(key)):
                errs.append(f"{label} مطلوب")
        clean["name"] = vd.clean_text(r.get("name"), 120)
        clean["title"] = vd.clean_text(r.get("title"), 120)
        clean["nationality"] = vd.clean_text(r.get("nationality"), 60) or "سعودي"
        clean["department"] = _s(r.get("department"))
        if clean["department"] and clean["department"] not in policy.DEPARTMENTS:
            match = next((d for d in policy.DEPARTMENTS if clean["department"] in d or d in clean["department"]), None)
            if match:
                clean["department"] = match
            else:
                errs.append(f"القسم «{clean['department']}» غير معروف")
        nid = re.sub(r"\D", "", _s(r.get("national_id")))
        if nid and not (len(nid) == 10 and nid[0] in "12"):
            errs.append("رقم الهوية/الإقامة لازم 10 أرقام ويبدأ بـ 1 أو 2")
        clean["national_id"] = nid
        for key, label in (("hire_date", "تاريخ المباشرة"), ("end_date", "تاريخ الانتهاء"), ("iqama_expiry", "انتهاء الإقامة")):
            try:
                clean[key] = _date(r.get(key))
                if clean[key] and not date(1970, 1, 1) <= clean[key] <= date(2100, 12, 31):
                    raise ValueError
            except (ValueError, OverflowError):
                errs.append(f"{label} غير صحيح")
                clean[key] = None
        for key, label in (("basic", "الأساسي"), ("housing", "السكن"), ("transport", "النقل"), ("leave_balance", "رصيد الإجازات")):
            try:
                clean[key] = _num(r.get(key))
            except ValueError:
                errs.append(f"{label} غير صحيح")
                clean[key] = None
        if clean.get("basic") == 0:
            errs.append("الأساسي لازم أكبر من صفر")
        ct = _s(r.get("contract_type"))
        clean["contract_type"] = "fixed" if ct in ("محدد", "محدد المدة", "fixed") else "open"
        if clean["contract_type"] == "fixed" and not clean.get("end_date"):
            errs.append("العقد المحدد يحتاج تاريخ انتهاء")
        for key, fn in (("iban", vd.iban), ("email", vd.email), ("phone", vd.phone)):
            try:
                clean[key] = fn(_s(r.get(key)))
            except ValueError as ex:
                errs.append(str(ex).rstrip("."))
                clean[key] = None
        clean["bank"] = vd.clean_text(r.get("bank"), 40) or None
        try:
            clean["code"] = vd.code(_s(r.get("code")), "الرقم الوظيفي")
            clean["manager_code"] = vd.code(_s(r.get("manager_code")), "رقم المدير")
            clean["attendance_id"] = vd.code(_s(r.get("attendance_id")), "رقم البصمة")
        except ValueError as ex:
            errs.append(str(ex).rstrip("."))
            clean.setdefault("code", None), clean.setdefault("manager_code", None), clean.setdefault("attendance_id", None)
        if clean["manager_code"] and clean["manager_code"] not in existing_codes | file_codes:
            errs.append(f"المدير {clean['manager_code']} غير موجود")
        clean["gosi_new_system"] = _s(r.get("gosi_new_system")) in ("نعم", "yes", "1", "true")
        role = _s(r.get("role")).lower() or "employee"
        if role not in ROLE_VALUES:
            errs.append("الدور غير معروف")
        clean["role"] = role
        for key, existing, label in (("code", existing_codes, "الرقم الوظيفي"), ("email", existing_emails, "البريد"),
                                     ("national_id", existing_ids, "رقم الهوية"), ("attendance_id", existing_att, "رقم البصمة")):
            v = clean.get(key)
            if v and (v in existing or v in seen[key]):
                errs.append(f"{label} مكرر")
            if v:
                seen[key].add(v)
        if clean.get("attendance_id") and str(clean["attendance_id"]).upper() in existing_codes | file_codes:
            errs.append("رقم البصمة يطابق رقمًا وظيفيًا لموظف آخر")
        if clean.get("code") and clean["code"] in existing_att:
            errs.append("الرقم الوظيفي يطابق رقم بصمة لموظف آخر")
        if clean.get("leave_balance") is not None and clean["leave_balance"] > 60:
            errs.append("رصيد الإجازات أكبر من المعقول (أكثر من 60)")
        results.append({"row": n, "data": clean, "errors": errs})
    return results


def _serialize(d):
    return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in d.items()}


def _deserialize(d):
    out = dict(d)
    for k in ("hire_date", "end_date", "iqama_expiry"):
        if out.get(k):
            out[k] = date.fromisoformat(out[k])
    return out


def _signer():
    from itsdangerous import URLSafeTimedSerializer
    return URLSafeTimedSerializer(current_app.config["SECRET_KEY"], salt="employee-import")


def pack(results, user_id):
    """الصفوف الصالحة تُوقّع وتُربط بالمستخدم، وصالحة ساعة وحدة فقط."""
    return _signer().dumps({"u": user_id, "rows": [_serialize(r["data"]) for r in results if not r["errors"]]})


def unpack(token, user_id):
    from itsdangerous import BadSignature, SignatureExpired
    try:
        data = _signer().loads(token, max_age=3600)
    except SignatureExpired:
        raise ValueError("انتهت مهلة المعاينة (ساعة). أعد رفع الملف.")
    except BadSignature:
        raise ValueError("بيانات الاستيراد غير صالحة. أعد رفع الملف.")
    if data.get("u") != user_id:
        raise ValueError("بيانات الاستيراد غير صالحة. أعد رفع الملف.")
    return [_deserialize(d) for d in data["rows"]]


def recheck_duplicates(records):
    """فحص ثاني لحظة الحفظ: لو أحد أضاف نفس الموظف بين المعاينة والحفظ، أو أُعيد إرسال نفس الملف مرتين."""
    codes = {e.code for e in Employee.query.all()}
    emails = {u.email for u in User.query.all() if u.email}
    ids = {e.national_id for e in Employee.query.all() if e.national_id}
    att = {str(e.attendance_id) for e in Employee.query.all() if e.attendance_id}
    keep, skipped = [], []
    for d in records:
        if (d.get("code") and d["code"] in codes) or (d.get("email") and d["email"] in emails) or \
                (d.get("national_id") and d["national_id"] in ids) or (d.get("attendance_id") and str(d["attendance_id"]) in att):
            skipped.append(d.get("name"))
            continue
        keep.append(d)
    return keep, skipped


def activation_code():
    alphabet = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"
    raw = "".join(secrets.choice(alphabet) for _ in range(8))
    return f"{raw[:4]}-{raw[4:]}"


SAFE_ROLES = {"employee", "manager"}


def import_records(records, actor_name, actor_is_admin=False):
    """غير مدير الموارد البشرية يقدر يستورد موظفين وmanagers فقط. الصلاحيات الأعلى تُعطى من صفحة الصلاحيات."""
    from .lifecycle import next_code
    created, codes = [], []
    for d in records:
        code = d.get("code") or next_code()
        basic = d["basic"]
        e = Employee(code=code, name=d["name"], email=d.get("email"), phone=d.get("phone"), nationality=d["nationality"],
                     is_saudi="سعود" in d["nationality"], gosi_new_system=d.get("gosi_new_system", False),
                     department=d["department"], title=d["title"], basic=basic,
                     housing=d["housing"] if d.get("housing") is not None else round(basic * policy.HOUSING_PCT),
                     transport=d["transport"] if d.get("transport") is not None else round(basic * policy.TRANSPORT_PCT),
                     hire_date=d["hire_date"], contract_type=d["contract_type"], end_date=d.get("end_date"),
                     iqama_expiry=d.get("iqama_expiry"), national_id=d.get("national_id"), bank=d.get("bank"),
                     iban=d.get("iban"), attendance_id=d.get("attendance_id"), contract_no=f"CT-{d['hire_date'].year}-{code.split('-')[-1]}")
        db.session.add(e)
        db.session.flush()
        if d.get("leave_balance") is not None:
            e.leave_used = max(0, e.entitlement - d["leave_balance"])
        created.append((e, d))
    by_code = {e.code: e for e, _ in created}
    for e, d in created:
        mc = d.get("manager_code")
        if mc:
            m = by_code.get(mc) or Employee.query.filter_by(code=mc).first()
            e.manager_id = m.id if m and m.id != e.id else None
        u = User(email=e.email, employee_id=e.id, active=True)
        role = d.get("role", "employee")
        u.role = role if (actor_is_admin or role in SAFE_ROLES) else "employee"
        db.session.add(u)
        db.session.flush()
        if e.email:
            raw = tokens.issue("set_password", u.id, 24 * 7)
            link = f"{current_app.config['BASE_URL']}/set-password/{raw}"
            mailer.queue(e.email, e.name, "تفعيل حسابك في مدار",
                         f"مرحبًا {e.first_name}،\n\nأصبح لك حساب في نظام مدار للموارد البشرية. من خلاله تقدّم طلباتك وتتابعها وتشوف رواتبك.\n"
                         "فعّل حسابك من الرابط (صالح 7 أيام).\n\nالموارد البشرية", employee_id=e.id, actions=[["تفعيل حسابي", link]], tag="تفعيل")
        else:
            code_raw = activation_code()
            tokens.issue_code("activate", u.id, code_raw, hours=24 * 14)
            codes.append({"code": e.code, "name": e.name, "activation": code_raw})
    audit.log("استيراد موظفين", f"{len(created)} موظف من ملف · بواسطة {actor_name}")
    return created, codes
