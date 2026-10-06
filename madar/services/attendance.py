"""استيراد سجلات البصمة ومطابقتها مع الطلبات.
أغلب أجهزة البصمة (مثل ZKTeco) تصدّر ملف CSV أو Excel فيه رقم الموظف ووقت البصمة. نتعرف على الأعمدة تلقائيًا.
للربط المباشر بالجهاز لاحقًا تقدر تستخدم مكتبة pyzk وتمرر نفس البيانات لدالة import_punches."""
import csv
import io
from datetime import date, datetime, time, timedelta

from flask import current_app

from .. import policy
from ..models import AttendanceDay, Employee, Request, db
from . import audit, mailer, settings, tasks

ID_COLS = ["code", "employee", "employee code", "emp", "emp no", "ac-no", "ac-no.", "ac no", "no", "no.", "userid",
           "user id", "user_id", "id", "رقم الموظف", "الرقم الوظيفي", "الرقم"]
TS_COLS = ["timestamp", "datetime", "date time", "date/time", "check time", "checktime", "punch time", "التاريخ والوقت", "وقت البصمة"]
DATE_COLS = ["date", "التاريخ", "اليوم"]
TIME_COLS = ["time", "الوقت", "الساعة"]
DT_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%d/%m/%Y %H:%M:%S",
              "%d/%m/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M"]


def _rows_from_file(fs):
    name = (fs.filename or "").lower()
    data = fs.read()
    if name.endswith(".xlsx"):
        from ..filesafe import check_zip
        check_zip(data)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = []
        for r in ws.iter_rows(values_only=True, max_col=30):
            rows.append(["" if c is None else c for c in r])
            if len(rows) > 200000:
                raise ValueError("الملف كبير جدًا. صدّر فترة أقصر من جهاز البصمة.")
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ValueError("ارفع ملف CSV أو Excel (xlsx) من جهاز البصمة.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("الملف فاضي.")
    return rows


def _find(header, names):
    h = [str(c).strip().lower() for c in header]
    for n in names:
        if n in h:
            return h.index(n)
    return None


def _parse_dt(value):
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    for f in DT_FORMATS:
        try:
            return datetime.strptime(s, f)
        except ValueError:
            continue
    return None


def parse_punches(fs):
    rows = _rows_from_file(fs)
    header, body = rows[0], rows[1:]
    i_id = _find(header, ID_COLS)
    i_ts, i_d, i_t = _find(header, TS_COLS), _find(header, DATE_COLS), _find(header, TIME_COLS)
    if i_ts is None and i_d is None and i_t is not None:
        i_ts, i_t = i_t, None  # مثل ZKTeco: عمود Time فيه التاريخ والوقت معًا
    if i_id is None or (i_ts is None and (i_d is None or i_t is None)):
        raise ValueError("ما قدرنا نتعرف على الأعمدة. لازم يكون فيه عمود لرقم الموظف (مثل AC-No أو Code) وعمود للوقت "
                         "(Timestamp) أو عمودين (Date و Time).")
    punches, bad = [], 0
    for r in body:
        try:
            code = str(r[i_id]).strip()
            if isinstance(r[i_id], float) and r[i_id].is_integer():
                code = str(int(r[i_id]))
            if i_ts is not None:
                dt = _parse_dt(r[i_ts])
            else:
                dv, tv = r[i_d], r[i_t]
                if isinstance(dv, datetime):
                    dv = dv.date()
                if isinstance(tv, datetime):
                    tv = tv.time()
                dt = datetime.combine(dv, tv) if isinstance(dv, date) and isinstance(tv, time) else _parse_dt(f"{dv} {tv}")
        except (IndexError, ValueError, TypeError):
            dt = None
        if not code or not dt:
            bad += 1
            continue
        punches.append((code, dt))
    return punches, bad


def _employee_map():
    """رقم البصمة له الأولوية. لو رقم ينطبق على أكثر من موظف (مثلًا رقم بصمة يساوي رقمًا وظيفيًا لشخص ثاني)
    نعتبره غامض ونتجاهله، بدل ما تنحسب بصمات شخص على شخص ثاني."""
    exact, loose = {}, {}
    for e in Employee.query.filter_by(status="active").all():
        k = e.code.upper()
        exact[k] = None if k in exact and exact[k] is not e else e
        if e.attendance_id:
            k = str(e.attendance_id).strip().upper()
            exact[k] = None if k in exact and exact[k] is not e else e
        suffix = e.code.split("-")[-1]
        loose[suffix] = None if suffix in loose and loose[suffix] is not e else e
    for k, e in loose.items():
        if k not in exact:
            exact[k] = e
        elif exact[k] is not None and e is not None and exact[k] is not e:
            exact[k] = None
    return exact


def import_punches(punches, skip_employee_id=None):
    """skip_employee_id: اللي رفع الملف ما تنحسب بصماته من ملفه (ما أحد يعدّل حضوره بنفسه)."""
    emap = _employee_map()
    days, unknown, skipped = {}, set(), 0
    lo_day, hi_day = policy.today() - timedelta(days=120), policy.today()
    for code, dt in punches:
        e = emap.get(code.strip().upper()[:40])
        if not e:
            unknown.add(code[:20])
            continue
        if e.id == skip_employee_id or not lo_day <= dt.date() <= hi_day:
            skipped += 1
            continue
        key = (e.id, dt.date())
        lo, hi = days.get(key, (dt.time(), dt.time()))
        days[key] = (min(lo, dt.time()), max(hi, dt.time()))
    for (emp_id, day), (cin, cout) in days.items():
        row = AttendanceDay.query.filter_by(employee_id=emp_id, day=day).first()
        if not row:
            row = AttendanceDay(employee_id=emp_id, day=day)
            db.session.add(row)
        row.check_in = cin if not row.check_in else min(row.check_in, cin)
        row.check_out = cout if cout != cin else row.check_out
    db.session.flush()
    tracked = {k[0] for k in days}
    if days:
        start, end = min(k[1] for k in days), max(k[1] for k in days)
        reconcile(start, end, tracked)
    return {"days": len(days), "employees": len(tracked), "unknown": sorted(unknown)[:20], "skipped": skipped}


def _approved_cover(emp, day):
    """هل يوجد طلب معتمد يغطي هذا اليوم؟ يرجع (نوع، ساعات الاستئذان)."""
    for r in Request.query.filter(Request.employee_id == emp.id, Request.status == "approved",
                                  Request.type.in_(["leave", "sick", "permission"])).all():
        d = r.data
        if r.type in ("leave", "sick"):
            start = date.fromisoformat(d["from"])
            if start <= day < start + timedelta(days=int(d["days"])):
                return r.type, None
        elif date.fromisoformat(d["date"]) == day:
            return "permission", float(d["hours"])
    return None, None


def reconcile(start, end, tracked_ids):
    """يحسب التأخير والغياب لكل يوم عمل، ويطابقه مع الإجازات والاستئذانات المعتمدة."""
    ws = datetime.strptime(settings.get("work_start"), "%H:%M").time()
    grace = settings.get_int("grace_minutes")
    weekend = settings.weekend_days()
    last = min(end, policy.today() - timedelta(days=1))
    day = start
    while day <= last:
        if day.weekday() not in weekend:
            for emp_id in tracked_ids:
                emp = db.session.get(Employee, emp_id)
                if not emp or emp.hire_date > day or emp.status != "active":
                    continue
                row = AttendanceDay.query.filter_by(employee_id=emp_id, day=day).first()
                cover, perm_hours = _approved_cover(emp, day)
                if row and row.check_in:
                    late = int((datetime.combine(day, row.check_in) - datetime.combine(day, ws)).total_seconds() // 60)
                    row.late_minutes = late if late > grace else 0
                    if row.status in ("excused", "deducted"):
                        continue
                    if row.late_minutes and cover == "permission" and perm_hours * 60 >= row.late_minutes:
                        row.status, row.note = "excused", "مغطى باستئذان معتمد"
                    else:
                        row.status = "late" if row.late_minutes else "present"
                else:
                    if not row:
                        row = AttendanceDay(employee_id=emp_id, day=day)
                        db.session.add(row)
                    if row.status in ("excused", "deducted"):
                        continue
                    if cover in ("leave", "sick"):
                        row.status, row.note = "excused", "إجازة معتمدة"
                    else:
                        row.status, row.note = "absent", "بدون بصمة ولا طلب"
        day += timedelta(days=1)
    db.session.flush()


def chase_absences():
    """يطالب الموظف برفع طلب عن كل غياب غير مبرر، ثم يبلّغ مديره إذا تأخر. يُشغّل يوميًا."""
    days_wait = settings.get_int("absence_reminder_days")
    base = current_app.config["BASE_URL"]
    done = []
    rows = AttendanceDay.query.filter(AttendanceDay.status == "absent",
                                      AttendanceDay.day <= policy.today() - timedelta(days=days_wait)).all()
    for a in rows:
        e = a.employee
        key = f"absence:{e.id}:{a.day.isoformat()}"
        if tasks.ensure(key, "manager", f"غياب بدون طلب: {e.name} يوم {a.day}",
                        "طُلب من الموظف رفع إجازة أو استئذان. إذا ما رفع، القرار لك: تبرير أو خصم",
                        employee_id=e.id, due=a.day + timedelta(days=days_wait + 2), link=f"/attendance?emp={e.id}",
                        assignee_id=e.manager.user.id if e.manager and e.manager.user else None):
            mailer.queue(e.email, e.name, f"غياب يوم {a.day} بدون طلب",
                         f"مرحبًا {e.first_name}،\n\nما فيه بصمة لك يوم {a.day} ولا يوجد طلب معتمد يغطيه.\n"
                         "إذا كنت في إجازة أو استئذان، ارفع الطلب الآن حتى لا يُحتسب غيابًا في الراتب.\n\nالموارد البشرية",
                         employee_id=e.id, actions=[["رفع طلب إجازة", f"{base}/requests/new/leave?from={a.day.isoformat()}"]],
                         tag="حضور")
            done.append(f"مطالبة {e.name} بتبرير غياب {a.day}")
    if done:
        audit.log("متابعة الغياب", f"{len(done)} يوم غياب بدون طلب", automated=True)
    return done


def month_summary(month, emp_ids=None):
    from .payroll import month_bounds
    start, end = month_bounds(month)
    q = AttendanceDay.query.filter(AttendanceDay.day.between(start, end))
    if emp_ids is not None:
        q = q.filter(AttendanceDay.employee_id.in_(emp_ids))
    out = {}
    for a in q.all():
        s = out.setdefault(a.employee_id, {"emp": a.employee, "present": 0, "late": 0, "late_minutes": 0, "absent": 0,
                                           "excused": 0, "deducted": 0, "days": {}})
        s[a.status if a.status in s else "present"] += 1
        s["late_minutes"] += a.late_minutes if a.status == "late" else 0
        s["days"][a.day.day] = a.status
    return out
